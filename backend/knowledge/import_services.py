import json
import re
from io import BytesIO

from django.db import transaction
from openpyxl import load_workbook
from pypdf import PdfReader

from audit.models import AdminAuditLog
from .models import Content,Discipline,DisciplineContent,KnowledgeStatus,Question,QuestionContentLink

MAX_XLSX_BYTES=10*1024*1024
MAX_PDF_BYTES=15*1024*1024
MAX_ROWS=2000
MAX_PDF_PAGES=300
MAX_PDF_TEXT=2_000_000

STATUS_ALIASES={
    "rascunho":"draft","draft":"draft","revisao":"review","revisão":"review","review":"review",
    "aprovado":"approved","approved":"approved","publicado":"published","published":"published",
    "inativo":"inactive","inactive":"inactive",
}
DIFFICULTY_ALIASES={
    "basica":"basic","básica":"basic","basic":"basic","facil":"basic","fácil":"basic",
    "intermediaria":"intermediate","intermediária":"intermediate","intermediate":"intermediate","media":"intermediate","média":"intermediate",
    "avancada":"advanced","avançada":"advanced","advanced":"advanced","dificil":"advanced","difícil":"advanced",
}
TYPE_ALIASES={
    "certo_errado":"certo_errado","certo/errado":"certo_errado","certo ou errado":"certo_errado","ce":"certo_errado",
    "multipla_escolha":"multipla_escolha","múltipla escolha":"multipla_escolha","multipla escolha":"multipla_escolha","me":"multipla_escolha",
}


def _clean(value):
    if value is None:return ""
    text=str(value).replace("\x00","").strip()
    return text


def _norm(value):
    return re.sub(r"\s+"," ",_clean(value)).strip().casefold()


def _bool(value,default=False):
    key=_norm(value)
    if key in {"1","true","sim","s","yes","y","x"}:return True
    if key in {"0","false","nao","não","n","no",""}:return False if value is not None else default
    raise ValueError("use SIM/NÃO ou TRUE/FALSE")


def _split_ids(value):
    if value is None or _clean(value)=="" :return []
    raw=re.split(r"[|;,]",_clean(value))
    values=[]
    for item in raw:
        item=item.strip()
        if not item:continue
        try:values.append(int(item))
        except ValueError:raise ValueError(f"ID inválido: {item}")
    return values


def _split_names(value):
    if value is None:return []
    return [item.strip() for item in re.split(r"[|;]",_clean(value)) if item.strip()]


def _safe_status(value):
    status=STATUS_ALIASES.get(_norm(value) or "draft")
    if not status:raise ValueError("status inválido")
    return status


def _safe_difficulty(value):
    difficulty=DIFFICULTY_ALIASES.get(_norm(value) or "intermediate")
    if not difficulty:raise ValueError("dificuldade inválida")
    return difficulty


def _safe_type(value):
    qtype=TYPE_ALIASES.get(_norm(value) or "certo_errado")
    if not qtype:raise ValueError("tipo inválido")
    return qtype


def _check_upload(upload,max_bytes,extensions):
    name=(getattr(upload,"name","") or "").lower()
    if not any(name.endswith(ext) for ext in extensions):raise ValueError(f"Formato inválido. Use {', '.join(extensions)}.")
    size=getattr(upload,"size",None)
    if size is not None and size>max_bytes:raise ValueError(f"Arquivo excede o limite de {max_bytes//1024//1024} MB.")


def _xlsx_rows(upload):
    _check_upload(upload,MAX_XLSX_BYTES,(".xlsx",))
    upload.seek(0)
    try:wb=load_workbook(filename=BytesIO(upload.read()),read_only=True,data_only=False)
    except Exception as exc:raise ValueError("Não foi possível abrir o XLSX. Verifique se o arquivo não está corrompido.") from exc
    ws=wb[wb.sheetnames[0]]
    if ws.max_row>MAX_ROWS+1:raise ValueError(f"A planilha possui mais de {MAX_ROWS} linhas de dados.")
    rows=list(ws.iter_rows(values_only=False))
    if not rows:return [],[]
    headers=[_norm(cell.value).replace(" ","_") for cell in rows[0]]
    if any(cell.data_type=="f" for row in rows for cell in row):
        raise ValueError("Planilhas com fórmulas não são aceitas. Cole apenas valores antes de importar.")
    data=[]
    for idx,row in enumerate(rows[1:],start=2):
        values=[cell.value for cell in row]
        if not any(_clean(value) for value in values):continue
        data.append((idx,{headers[i]:values[i] if i<len(values) else None for i in range(len(headers)) if headers[i]}))
    return headers,data


def _content_ids(row):
    ids=_split_ids(row.get("conteudo_ids"))
    if ids:
        existing={item.id:item for item in Content.objects.filter(id__in=ids)}
        missing=[item for item in ids if item not in existing]
        if missing:raise ValueError("conteudo_ids inexistente(s): "+", ".join(map(str,missing)))
        return ids
    names=_split_names(row.get("conteudos"))
    resolved=[]
    for name in names:
        matches=list(Content.objects.filter(title__iexact=name).values_list("id",flat=True)[:2])
        if not matches:raise ValueError(f"conteúdo não encontrado: {name}")
        if len(matches)>1:raise ValueError(f"título de conteúdo ambíguo: {name}; use conteudo_ids")
        resolved.append(matches[0])
    return resolved


def _discipline_ids(row):
    ids=_split_ids(row.get("disciplina_ids"))
    if ids:
        existing=set(Discipline.objects.filter(id__in=ids).values_list("id",flat=True))
        missing=[item for item in ids if item not in existing]
        if missing:raise ValueError("disciplina_ids inexistente(s): "+", ".join(map(str,missing)))
        return ids
    short_names=_split_names(row.get("disciplinas"))
    resolved=[]
    for short in short_names:
        match=Discipline.objects.filter(short_name__iexact=short).first()
        if not match:raise ValueError(f"sigla de disciplina não encontrada: {short}")
        resolved.append(match.id)
    return resolved


def parse_question_xlsx(upload):
    headers,rows=_xlsx_rows(upload)
    required={"enunciado","tipo","resposta"}
    missing=sorted(required-set(headers))
    if missing:raise ValueError("Colunas obrigatórias ausentes: "+", ".join(missing))
    result=[]
    for row_number,row in rows:
        errors=[];warnings=[]
        try:
            statement=_clean(row.get("enunciado"))
            if len(statement)<12:raise ValueError("enunciado deve ter ao menos 12 caracteres")
            qtype=_safe_type(row.get("tipo"))
            options=[]
            for letter in "abcde":
                value=_clean(row.get(f"alternativa_{letter}"))
                if value:options.append(value)
            raw_answer=_clean(row.get("resposta"))
            if qtype=="certo_errado":
                key=_norm(raw_answer)
                if key in {"certo","true","verdadeiro","1"}:answer=True
                elif key in {"errado","false","falso","0"}:answer=False
                else:raise ValueError("resposta de Certo/Errado deve ser CERTO ou ERRADO")
                options=[]
            else:
                if len(options)<2:raise ValueError("múltipla escolha exige ao menos alternativa_a e alternativa_b")
                key=raw_answer.strip().upper().replace(")","").replace(".","")
                if key in {"A","B","C","D","E"}:
                    position=ord(key)-65
                    if position>=len(options):raise ValueError("resposta aponta para alternativa vazia")
                    answer=options[position]
                elif raw_answer in options:answer=raw_answer
                else:raise ValueError("resposta deve ser A/B/C/D/E ou o texto exato de uma alternativa")
            year=None
            if _clean(row.get("ano")):
                year=int(float(row.get("ano")))
                if year<1900 or year>2100:raise ValueError("ano fora do intervalo 1900-2100")
            content_ids=_content_ids(row)
            duplicate=Question.objects.filter(statement__iexact=statement,question_type=qtype).exists()
            if duplicate:warnings.append("questão idêntica já existe e será ignorada")
            result.append({"row":row_number,"valid":not duplicate,"skip":duplicate,"errors":[],"warnings":warnings,"data":{
                "statement":statement,"question_type":qtype,"options_json":options,"answer_json":answer,
                "explanation":_clean(row.get("comentario")),"difficulty":_safe_difficulty(row.get("dificuldade")),
                "source":_clean(row.get("fonte"))[:240],"banca":_clean(row.get("banca"))[:120],"year":year,
                "status":_safe_status(row.get("status")),"requires_review":_bool(row.get("exigir_revisao")),
                "content_ids":content_ids,
            }})
        except Exception as exc:
            errors.append(str(exc))
            result.append({"row":row_number,"valid":False,"skip":False,"errors":errors,"warnings":warnings,"data":{"statement":_clean(row.get("enunciado"))}})
    return result


def apply_question_import(rows,user):
    created=[]
    with transaction.atomic():
        for item in rows:
            if not item["valid"] or item["skip"]:continue
            data=item["data"];content_ids=data.pop("content_ids")
            question=Question.objects.create(created_by=user,updated_by=user,**data)
            for content_id in content_ids:QuestionContentLink.objects.get_or_create(question=question,content_id=content_id,defaults={"linked_by":user})
            created.append(question.id)
        AdminAuditLog.objects.create(actor=user,action="IMPORTACAO_QUESTOES_XLSX",detail=f"{len(created)} questão(ões) importada(s) pela biblioteca ROOT.")
    return created


def parse_content_xlsx(upload):
    headers,rows=_xlsx_rows(upload)
    if "titulo" not in headers:raise ValueError("Coluna obrigatória ausente: titulo")
    result=[]
    for row_number,row in rows:
        errors=[];warnings=[]
        try:
            title=_clean(row.get("titulo"))
            if len(title)<3:raise ValueError("título deve ter ao menos 3 caracteres")
            discipline_ids=_discipline_ids(row)
            duplicate=Content.objects.filter(title__iexact=title).exists()
            if duplicate:warnings.append("conteúdo com este título já existe e será ignorado")
            notice=_norm(row.get("aviso"))
            notice_kind={"novo":"new","new":"new","atualizado":"updated","updated":"updated","":"None","nenhum":"None"}.get(notice)
            if notice_kind is None:raise ValueError("aviso deve ser NENHUM, NOVO ou ATUALIZADO")
            result.append({"row":row_number,"valid":not duplicate,"skip":duplicate,"errors":[],"warnings":warnings,"data":{
                "title":title,"objective":_clean(row.get("objetivo")),"description":_clean(row.get("descricao")),
                "card_text":_clean(row.get("resumo_card")),"body":_clean(row.get("corpo")),
                "cover_image_url":_clean(row.get("capa_url"))[:2048],"video_url":_clean(row.get("video_url"))[:2048],
                "video_label":_clean(row.get("video_rotulo"))[:160],"material_url":_clean(row.get("material_url"))[:2048],
                "material_label":_clean(row.get("material_rotulo"))[:160],"notice_kind":None if notice_kind=="None" else notice_kind,
                "status":_safe_status(row.get("status")),"requires_review":_bool(row.get("exigir_revisao")),
                "discipline_ids":discipline_ids,
            }})
        except Exception as exc:
            errors.append(str(exc))
            result.append({"row":row_number,"valid":False,"skip":False,"errors":errors,"warnings":warnings,"data":{"title":_clean(row.get("titulo"))}})
    return result


def apply_content_import(rows,user):
    created=[]
    with transaction.atomic():
        for item in rows:
            if not item["valid"] or item["skip"]:continue
            data=item["data"];discipline_ids=data.pop("discipline_ids")
            content=Content.objects.create(created_by=user,updated_by=user,**data)
            for discipline_id in discipline_ids:DisciplineContent.objects.get_or_create(discipline_id=discipline_id,content=content,defaults={"linked_by":user})
            created.append(content.id)
        AdminAuditLog.objects.create(actor=user,action="IMPORTACAO_CONTEUDOS_XLSX",detail=f"{len(created)} conteúdo(s) importado(s) pela biblioteca ROOT.")
    return created


def pdf_to_content(upload,metadata):
    _check_upload(upload,MAX_PDF_BYTES,(".pdf",))
    upload.seek(0)
    try:reader=PdfReader(BytesIO(upload.read()))
    except Exception as exc:raise ValueError("Não foi possível abrir o PDF.") from exc
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):raise ValueError("PDF protegido por senha não pode ser importado.")
        except Exception as exc:raise ValueError("PDF protegido por senha não pode ser importado.") from exc
    if len(reader.pages)>MAX_PDF_PAGES:raise ValueError(f"PDF excede o limite de {MAX_PDF_PAGES} páginas.")
    chunks=[]
    total=0
    for index,page in enumerate(reader.pages,start=1):
        text=(page.extract_text() or "").replace("\x00","").strip()
        if not text:continue
        chunk=f"Página {index}\n\n{text}"
        total+=len(chunk)
        if total>MAX_PDF_TEXT:raise ValueError("Texto extraído do PDF excede o limite permitido.")
        chunks.append(chunk)
    body="\n\n".join(chunks).strip()
    if len(body)<50:raise ValueError("O PDF não possui texto extraível suficiente. PDFs escaneados precisam de OCR antes da importação.")
    title=_clean(metadata.get("title")) or re.sub(r"\.pdf$","",getattr(upload,"name","conteudo"),flags=re.I)
    if len(title)<3:raise ValueError("Informe um título válido.")
    discipline_ids=_split_ids(metadata.get("disciplineIds"))
    missing=set(discipline_ids)-set(Discipline.objects.filter(id__in=discipline_ids).values_list("id",flat=True))
    if missing:raise ValueError("Disciplina(s) inexistente(s): "+", ".join(map(str,sorted(missing))))
    duplicate=Content.objects.filter(title__iexact=title).exists()
    return {
        "title":title[:220],"objective":_clean(metadata.get("objective")),"description":_clean(metadata.get("description")),
        "card_text":_clean(metadata.get("cardText")),"body":body,"status":_safe_status(metadata.get("status")),
        "requires_review":_bool(metadata.get("requiresReview")),"discipline_ids":discipline_ids,"duplicate":duplicate,
        "pages":len(reader.pages),"characters":len(body),
    }


def apply_pdf_content(data,user):
    if data.get("duplicate"):return None
    payload={key:value for key,value in data.items() if key not in {"discipline_ids","duplicate","pages","characters"}}
    discipline_ids=data["discipline_ids"]
    with transaction.atomic():
        content=Content.objects.create(created_by=user,updated_by=user,**payload)
        for discipline_id in discipline_ids:DisciplineContent.objects.get_or_create(discipline_id=discipline_id,content=content,defaults={"linked_by":user})
        AdminAuditLog.objects.create(actor=user,action="IMPORTACAO_CONTEUDO_PDF",detail=f"Conteúdo #{content.id} importado do PDF pela biblioteca ROOT.")
    return content.id


def import_summary(rows):
    valid=sum(1 for item in rows if item["valid"] and not item["skip"])
    skipped=sum(1 for item in rows if item["skip"])
    invalid=sum(1 for item in rows if not item["valid"] and not item["skip"])
    preview=[]
    for item in rows[:100]:
        data=item.get("data") or {}
        preview.append({
            "row":item["row"],"status":"skip" if item["skip"] else "valid" if item["valid"] else "invalid",
            "label":data.get("statement") or data.get("title") or "",
            "errors":item["errors"],"warnings":item["warnings"],
        })
    return {"validRows":valid,"skippedRows":skipped,"invalidRows":invalid,"preview":preview}
